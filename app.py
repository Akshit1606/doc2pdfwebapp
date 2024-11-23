import os
from flask import Flask, render_template, request, send_from_directory, jsonify
from werkzeug.utils import secure_filename
from docx import Document
from PyPDF2 import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import logging

# Initialize Flask app
app = Flask(__name__)

# Configure environment variables
port = int(os.environ.get("PORT", 5000))

# Configure upload folder and allowed extensions
UPLOAD_FOLDER = '/tmp/uploads'  # Temporary storage for uploaded files
OUTPUT_FOLDER = '/tmp/converted_pdfs'  # Temporary storage for converted PDFs
ALLOWED_EXTENSIONS = {'docx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['OUTPUT_FOLDER'] = OUTPUT_FOLDER

# Ensure the upload and output directories exist
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# Configure logging
logging.basicConfig(level=logging.INFO)

# Check file extension
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

# Function to extract metadata from the docx file
def get_docx_metadata(file_path):
    try:
        document = Document(file_path)
        metadata = {
            'title': document.core_properties.title,
            'author': document.core_properties.author,
            'subject': document.core_properties.subject,
            'keywords': document.core_properties.keywords,
            'created': document.core_properties.created,
        }
        logging.debug(f"Metadata extracted: {metadata}")
        return metadata
    except Exception as e:
        logging.error(f"Error extracting metadata: {str(e)}")
        return None

# Convert DOCX to PDF using ReportLab
def convert_docx_to_pdf(docx_path, pdf_path):
    doc = Document(docx_path)
    c = canvas.Canvas(pdf_path, pagesize=letter)
    width, height = letter
    y_position = height - 40
    for para in doc.paragraphs:
        c.drawString(40, y_position, para.text)
        y_position -= 12
        if y_position < 40:
            c.showPage()
            y_position = height - 40
    c.save()
    logging.debug(f"Converted DOCX to PDF: {pdf_path}")
    return pdf_path

# Home page route
@app.route('/')
def index():
    return render_template('index.html')

# File upload route
@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"})
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"})
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        uploaded_file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(uploaded_file_path)
        
        # Extract metadata from DOCX file
        metadata = get_docx_metadata(uploaded_file_path)
        
        # Convert DOCX to PDF
        pdf_filename = f"{os.path.splitext(filename)[0]}.pdf"
        pdf_file_path = os.path.join(app.config['OUTPUT_FOLDER'], pdf_filename)
        convert_docx_to_pdf(uploaded_file_path, pdf_file_path)
        
        # Return metadata and download link
        return render_template('result.html', metadata=metadata, pdf_filename=pdf_filename)

# Download route for converted PDF
@app.route('/download/<filename>')
def download_file(filename):
    return send_from_directory(app.config['OUTPUT_FOLDER'], filename)

# Error handling route
@app.errorhandler(404)
def not_found_error(error):
    return render_template('404.html'), 404

# Run the app
if __name__ == '__main__':
    app.run(host="0.0.0.0", port=port)
